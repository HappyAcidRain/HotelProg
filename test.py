from docx import Document

# создание пустого документа
doc = Document()
# добавляем пустую таблицу 2х2 ячейки
table = doc.add_table(rows=2, cols=2)

# выбор стиля таблицы
table.style = 'Light Shading Accent 1'

# выбор ячейки
cell = table.cell(0, 1)

# запись в ячейку данных
cell.text = 'hello world'


# функция внесения значений
def ins(x, y, z):
    c = table.cell(x, y)
    c.text = z


# функция внесения имен таблицы
def cName(y, z):
    c = table.cell(0, y)
    c.text = z

cName(0, "имя 1 ")

ins(1, 0, 'test 2def')

ins(1, 1, 'test 4def')

# сохранение изменений
doc.save('test.docx')
